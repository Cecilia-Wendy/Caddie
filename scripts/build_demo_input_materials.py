#!/usr/bin/env python3
"""Build a realistic raw-material pack for testing Caddie's experience intake."""

from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import csv

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "caddie_demo_input_materials"
FINAL = OUT / "林知夏_Caddie经历输入测试材料包"
INK = RGBColor(31, 41, 55)
BLUE = RGBColor(38, 92, 142)
MUTED = RGBColor(100, 116, 139)
NOTICE = "全部人物、机构、项目与数字均为虚构，仅用于 Caddie 产品测试。"


def set_font(run, size=10.5, bold=False, color=INK):
    run.font.name = "STHeiti"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "STHeiti")
    fonts.set(qn("w:hAnsi"), "STHeiti")
    fonts.set(qn("w:eastAsia"), "STHeiti")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_doc(doc, running_label):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    normal = doc.styles["Normal"]
    normal.font.name = "STHeiti"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "STHeiti")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "STHeiti")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before, after in (
        ("Heading 1", 16, 15, 7),
        ("Heading 2", 12.5, 11, 5),
        ("Heading 3", 11, 8, 3),
    ):
        style = doc.styles[name]
        style.font.name = "STHeiti"
        style._element.rPr.rFonts.set(qn("w:ascii"), "STHeiti")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "STHeiti")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    hp = section.header.paragraphs[0]
    hp.text = running_label
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(hp.runs[0], 8.5, color=MUTED)
    fp = section.footer.paragraphs[0]
    fp.text = NOTICE
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(fp.runs[0], 8, color=MUTED)


def add_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, 23, True, INK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(14)
    r2 = p2.add_run(subtitle)
    set_font(r2, 10.5, color=MUTED)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r)


def build_resume():
    doc = Document()
    configure_doc(doc, "原始材料 · 01 / 个人简历")
    add_title(doc, "林知夏", "复旦大学管理科学与工程硕士｜求职方向：产品经理 / 商业分析")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("上海｜138-0000-2468｜zhixia.lin@example.test｜2027 届")
    set_font(r, 10, color=MUTED)

    doc.add_heading("教育背景", level=1)
    for school, degree, dates, detail in [
        ("复旦大学", "管理科学与工程 硕士", "2024.09–2027.06", "GPA 3.72/4.0；一等奖学金；课程：机器学习、因果推断、运营管理"),
        ("中山大学", "信息管理与信息系统 本科", "2020.09–2024.06", "GPA 3.78/4.0；国家奖学金；优秀毕业生"),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(school + "  "), 11, True)
        set_font(p.add_run(degree), 10.5)
        set_font(p.add_run("  " + dates), 9.5, color=MUTED)
        add_bullet(doc, detail)

    doc.add_heading("实习经历", level=1)
    exp_data = [
        ("星河智能科技（虚构）", "AI 产品经理实习生", "2026.03–2026.07",
         [
             "参与企业知识助手从需求调研到灰度上线，访谈售前和产品专家，整理高频任务。",
             "协助搭建 AI 回答评测集，跟进灰度反馈和问题回流；试点采纳率约 70%。",
             "推动答案引用和低置信度提示等功能迭代。",
         ]),
        ("远见咨询（虚构）", "数字化战略咨询实习生", "2025.07–2025.11",
         [
             "参与连锁零售会员增长项目，负责数据清洗、用户分群和部分客户汇报材料。",
             "支持制造业 CRM 蓝图项目，完成访谈纪要与流程梳理。",
         ]),
        ("青禾消费（虚构）", "商业分析实习生", "2024.12–2025.04",
         [
             "分析新品上市后的门店动销、库存和活动表现，支持渠道调整。",
             "搭建周度看板，跟踪区域和门店层级指标。",
         ]),
    ]
    for company, role, dates, bullets in exp_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(company + "  "), 11, True)
        set_font(p.add_run(role), 10.5)
        set_font(p.add_run("  " + dates), 9.5, color=MUTED)
        for item in bullets:
            add_bullet(doc, item)

    doc.add_heading("项目与校园经历", level=1)
    p = doc.add_paragraph()
    set_font(p.add_run("校园闲置物品平台「拾光」｜联合创始人 / 产品负责人  "), 11, True)
    set_font(p.add_run("2023.03–2024.05"), 9.5, color=MUTED)
    add_bullet(doc, "通过访谈设计校园认证、信用记录和集中交付点；累计注册用户约 2,000 人。")
    add_bullet(doc, "负责需求、原型和运营，开发由技术合伙人完成。")

    doc.add_heading("技能", level=1)
    p = doc.add_paragraph()
    set_font(p.add_run("SQL、Python、Tableau、Figma、Axure、A/B 测试、用户研究、PRD；英语 CET-6 612"))
    path = FINAL / "01_第一轮导入" / "林知夏_原始简历.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def build_project_notes():
    doc = Document()
    configure_doc(doc, "原始材料 · 02 / 工作笔记")
    add_title(doc, "企业知识助手项目｜零散工作笔记", "来自周报、会议记录和个人备忘的混合整理；部分口径尚未确认")
    doc.add_heading("3 月 18 日｜需求访谈", level=1)
    add_bullet(doc, "售前找资料很慢，有人说平均十几分钟，但没做正式测量。")
    add_bullet(doc, "大家最担心 AI 编答案，所以可能要显示来源；权限问题也有人提。")
    add_bullet(doc, "访谈人数：大概 10 多个售前，另有几位产品专家。具体名单要找会议记录。")

    doc.add_heading("4 月 2 日｜方案讨论", level=1)
    add_bullet(doc, "我画了“问题—答案—引用”的原型；算法同学建议增加拒答。")
    add_bullet(doc, "RAG / 微调为什么这样选，当时主要听算法同学判断，我负责从产品风险角度提要求。")
    add_bullet(doc, "版本一期先不做多轮任务，只覆盖资料查询。")

    doc.add_heading("5 月 9 日｜评测", level=1)
    add_bullet(doc, "评测题从历史咨询问题和专家补充中来，最后好像 180 道。")
    add_bullet(doc, "最初只有正确 / 错误，后来拆成正确性、完整性、引用、拒答、表达。")
    add_bullet(doc, "有一版整体正确率不错，但引用并不支持答案——这是比较大的坑。")

    doc.add_heading("6 月 24 日｜灰度复盘", level=1)
    add_bullet(doc, "30 人左右参与，两周总查询 486 次。")
    add_bullet(doc, "采纳率一份表里是 71%，周报里写过 70%；需要确认分母是否只算有效回答。")
    add_bullet(doc, "找资料时间从 18 分钟降到约 5 分钟，但 18 分钟来自访谈回忆，不是系统日志。")
    add_bullet(doc, "下一步：知识缺口回流、引用定位、权限标签。")

    doc.add_heading("我到底负责了什么（未整理版）", level=1)
    add_bullet(doc, "我：访谈、需求归纳、原型、评测规则、灰度运营、复盘。")
    add_bullet(doc, "算法：检索、重排、模型选择、工程实现。")
    add_bullet(doc, "共同：评测题、版本验收、问题分类。")
    add_bullet(doc, "不能写成“独立搭建 RAG 系统”。")

    doc.add_heading("还没想清楚的问题", level=1)
    add_bullet(doc, "如果被问为什么不用微调，我怎么讲才不装懂？")
    add_bullet(doc, "71% 到底算不算好？覆盖率和采纳率怎么一起讲？")
    add_bullet(doc, "最失败的决策是什么：第一版评测指标，还是范围定得太宽？")
    path = FINAL / "02_第二轮补充" / "企业知识助手_零散工作笔记.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def write_text_files():
    files = {}
    files["meeting"] = FINAL / "02_第二轮补充" / "企业知识助手_灰度复盘会议纪要.txt"
    files["meeting"].write_text(
        """会议：企业知识助手灰度复盘
日期：2026-06-24
参会：产品、算法、售前试点成员（虚构）

1. 本轮 30 名试点成员，统计区间 6/10—6/23。
2. 日志显示总查询 486 次，其中 421 次产生有效回答。
3. “有帮助”或复制答案共 299 次。会上有人直接用 299/421 得到 71.0%。
4. 如果按全部查询计算，则为 61.5%。后续对外表达必须同时给出口径。
5. 随机访谈 8 人，自报单次找资料时间由约 18 分钟降为 4—6 分钟；不是全量日志指标。
6. 主要问题：
   - 23 次无答案集中在新品和地区政策；
   - 12 条答案的引用不能完整支持结论；
   - 权限标签缺失导致 4 条内容不应对该试点组可见（上线前已拦截）。
7. 后续动作：补知识、上线引用定位、增加权限检查、把引用一致性设为发布门槛。

职责记录：
- 林知夏：主持复盘、整理日志、定义产品指标与后续需求。
- 周工 / 赵工：算法和工程问题定位。
- 售前专家：确认业务答案与资料有效性。

""" + NOTICE + "\n",
        encoding="utf-8",
    )
    files["transcript"] = FINAL / "03_第三轮验证" / "模拟面试_企业知识助手追问逐字稿.txt"
    files["transcript"].parent.mkdir(parents=True, exist_ok=True)
    files["transcript"].write_text(
        """面试官：你说采纳率提升到 71%，这个口径是什么？
林知夏：用户觉得有帮助的比例，大概是 71%。

面试官：分母是所有问题还是有效回答？
林知夏：这里我需要确认。印象中是有效回答，因为无答案问题没有参与评价。

面试官：为什么选择 RAG 而不是微调？
林知夏：因为我们的知识更新比较快，还要展示来源。技术选型是算法同学主导，我当时重点定义的是来源可信、低置信度拒答和权限这几个产品要求。

面试官：这个项目最失败的地方是什么？
林知夏：第一版评测只看答案是否正确，后来发现引用并不能支持答案。这个问题是灰度前抽查发现的，我们增加引用一致性指标才解决。

面试官：这个错误为什么一开始没想到？
林知夏：我当时把“给出引用”等同于“引用可信”，对企业场景的风险理解不够深。现在我会在需求阶段就把业务风险转成独立验收指标。

""" + NOTICE + "\n",
        encoding="utf-8",
    )
    files["boundary"] = FINAL / "03_第三轮验证" / "本人贡献边界_给Caddie的事实说明.md"
    files["boundary"].write_text(
        """# 本人贡献边界

## 已确认

- 独立完成：用户访谈提纲、需求归类、产品原型、灰度反馈流程、复盘材料。
- 主导但非独立完成：评测规则与版本验收；评测题由本人、算法和业务专家共同完成。
- 未负责：检索算法实现、模型训练、服务部署、底层数据管道。

## 数字口径

- 试点 30 人、两周查询 486 次。
- 421 次产生有效回答；299 次被标记“有帮助”或被复制。
- 71% = 299 / 421；若分母为全部查询，则为 61.5%。
- “18 分钟降到 4—6 分钟”来自 8 人访谈，不应表述为全量系统统计。

## 暂未确认

- 访谈总人数到底是 16 人还是 17 人。
- 项目开始日是 3 月 11 日还是 3 月 14 日。
- 引用一致性问题是 12 条还是 13 条，其中一条是否为重复记录。

> """ + NOTICE + "\n",
        encoding="utf-8",
    )
    files["readme"] = FINAL / "00_使用说明_按这个顺序跑一遍.md"
    files["readme"].write_text(
        """# Caddie 经历输入测试脚本

虚拟求职者：林知夏。建议从一个空白测试库开始。

## 第一轮：只给简历

上传 `01_第一轮导入/林知夏_原始简历.docx`。

观察：

1. 是否识别出教育、3 段实习和 1 段创业经历；
2. 是否把“企业知识助手”与“AI 回答评测”先视为同一经历下的潜在项目；
3. 是否避免把“参与、协助”自动夸大成“主导、独立完成”；
4. 是否主动追问数字口径和本人职责。

## 第二轮：补原始证据

依次上传：

1. `02_第二轮补充/企业知识助手_零散工作笔记.docx`
2. `02_第二轮补充/企业知识助手_灰度复盘会议纪要.txt`
3. `02_第二轮补充/企业知识助手_灰度数据.csv`
4. `02_第二轮补充/林知夏_实习项目清单.csv`

观察：

- 是否发现 71% 与 61.5% 是不同分母；
- 是否把 18 分钟到 4—6 分钟标记为访谈样本，而非全量日志；
- 是否拆出“企业知识助手 MVP”和“AI 回答质量评测体系”两个相关项目；
- 是否记录算法实现不属于林知夏的贡献。

## 第三轮：验证追问与纠错

上传：

1. `03_第三轮验证/本人贡献边界_给Caddie的事实说明.md`
2. `03_第三轮验证/模拟面试_企业知识助手追问逐字稿.txt`

然后让 Caddie：

1. 生成结构化项目档案；
2. 列出仍需本人确认的事实；
3. 生成 8 个高风险追问；
4. 给出一版不夸大贡献的 STAR 回答；
5. 检查生成结果是否错误采用了 71% 的模糊口径。

## 建议验收标准

- 不编造未提供的数字、客户名称或技术实现；
- 区分本人完成、主导协作、团队完成与未负责；
- 冲突数字不擅自选一个，而是保留口径说明；
- 能把原始材料关联回经历、项目、知识、追问和证据；
- 仍不确定的日期、人数和重复记录进入待确认清单。

> """ + NOTICE + "\n",
        encoding="utf-8",
    )
    return list(files.values())


def write_csv_files():
    path = FINAL / "02_第二轮补充" / "企业知识助手_灰度数据.csv"
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["指标", "数值", "分母/样本", "来源", "备注"])
        writer.writerows([
            ["试点人数", 30, "全部试点成员", "试点名单", "两周灰度"],
            ["查询次数", 486, "全部查询", "产品日志", "含无答案"],
            ["有效回答", 421, "全部查询", "产品日志", "有效回答覆盖率 86.6%"],
            ["有帮助或复制", 299, "有效回答", "行为日志", "299/421=71.0%"],
            ["有帮助或复制", 299, "全部查询", "行为日志", "299/486=61.5%"],
            ["平均原查找时间", "约18分钟", "访谈8人", "用户回忆", "非全量日志"],
            ["灰度后查找时间", "4—6分钟", "访谈8人", "用户回忆", "区间值"],
            ["引用不一致", 12, "抽查记录", "人工复核", "可能有1条重复，待确认"],
        ])
    path2 = FINAL / "02_第二轮补充" / "林知夏_实习项目清单.csv"
    with path2.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["经历", "可能的项目", "时间", "我的角色", "结果/产出", "待补信息"])
        writer.writerows([
            ["星河智能科技", "企业知识助手MVP", "2026.03-06", "产品实习生", "完成灰度", "项目准确起始日"],
            ["星河智能科技", "AI回答评测体系", "2026.04-06", "产品与评测协作", "180题评测集", "具体个人完成题数"],
            ["远见咨询", "连锁零售会员增长", "2025.07-09", "数据分析与材料支持", "20店试点复购率+4.2pp", "本人是否参与试点执行"],
            ["远见咨询", "制造业CRM蓝图", "2025.09-11", "访谈与流程梳理", "12项一期需求", "需求优先级方法"],
            ["青禾消费", "新品上市渠道诊断", "2025.01-03", "商业分析实习生", "试点单店周销+18%", "对照口径"],
            ["校园创业", "拾光闲置平台", "2023.03-2024.05", "联合创始人/产品", "注册约2300人", "停止运营原因"],
        ])
    return [path, path2]


def make_zip():
    zip_path = OUT / "林知夏_Caddie经历输入测试材料包.zip"
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
        for path in sorted(FINAL.rglob("*")):
            if path.is_file():
                zf.write(path, path.relative_to(FINAL.parent))
    return zip_path


def main():
    FINAL.mkdir(parents=True, exist_ok=True)
    resume = build_resume()
    notes = build_project_notes()
    write_text_files()
    write_csv_files()
    zip_path = make_zip()
    print(resume)
    print(notes)
    print(zip_path)


if __name__ == "__main__":
    main()
